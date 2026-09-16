"""Generate Milestone 2 sample project documents."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent / "milestone2"
ROOT.mkdir(parents=True, exist_ok=True)


def write_text_pdf(path: Path, lines: list[str], lines_per_page: int = 38) -> None:
    pages = [lines[i : i + lines_per_page] for i in range(0, len(lines), lines_per_page)] or [[]]
    objects: list[bytes] = []
    page_ids = []
    content_ids = []

    # We assemble a simple PDF: catalog, pages tree, pages, contents, font.
    # Object numbers are assigned after we know page count.
    font_obj = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    page_streams = []
    for page_lines in pages:
        escaped = []
        for line in page_lines:
            text = (
                line.replace("\\", "\\\\")
                .replace("(", "\\(")
                .replace(")", "\\)")
            )
            escaped.append(text)
        ops = "BT /F1 11 Tf 50 740 Td\n"
        for text in escaped:
            ops += f"({text}) Tj\n0 -18 Td\n"
        ops += "ET\n"
        page_streams.append(ops.encode("latin-1", errors="replace"))

    # object 1 catalog, 2 pages, then pairs of page+content, then font
    n_pages = len(page_streams)
    first_page_id = 3
    font_id = first_page_id + n_pages * 2

    kids = []
    page_objects = []
    for index, stream in enumerate(page_streams):
        page_id = first_page_id + index * 2
        content_id = page_id + 1
        kids.append(f"{page_id} 0 R")
        page_objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        page_objects.append(
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        )

    catalog = b"<< /Type /Catalog /Pages 2 0 R >>"
    pages_obj = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {n_pages} >>".encode()
    all_objects = [catalog, pages_obj, *page_objects, font_obj]

    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(all_objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode("ascii"))
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")
    xref_start = len(pdf)
    pdf.extend(f"xref\n0 {len(all_objects) + 1}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        (
            f"trailer << /Size {len(all_objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_start}\n%%EOF\n"
        ).encode("ascii")
    )
    path.write_bytes(bytes(pdf))


SOFTWARE_PLAN = [
    "PROJECT NAME: Campus Attendance Platform",
    "OBJECTIVE: Build a QR-based attendance system for the computer science department.",
    "SCOPE: In scope: student mobile check-in, faculty dashboard, and attendance reports. Out of scope: payroll and hostel management.",
    "START DATE: 2026-01-15",
    "DEADLINE: 2026-06-30",
    "",
    "TEAM MEMBERS",
    "Aisha Khan: Project manager",
    "Priya Sharma: Mobile developer",
    "Daniel Okonkwo: Backend developer",
    "Ravi Patel: QA intern",
    "",
    "DELIVERABLES",
    "Deliverable | Owner | Deadline | Status",
    "Mobile check-in app | Priya Sharma | 2026-03-15 | In progress",
    "Faculty dashboard | Daniel Okonkwo | 2026-04-10 | Not started",
    "Attendance reports | Priya Sharma | 2026-05-20 | Not started",
    "",
    "MILESTONES",
    "Milestone | Date | Owner | Status",
    "Requirements freeze | 2026-02-01 | Aisha Khan | Completed",
    "Beta release | 2026-05-01 | Priya Sharma | At risk",
    "Final delivery | 2026-06-30 | Aisha Khan | Not started",
    "",
    "TASKS",
    "Task | Owner | Deadline | Status | Depends On",
    "Design API | Daniel Okonkwo | 2026-02-20 | Delayed | Requirements freeze",
    "Implement QR scan | Priya Sharma | 2026-03-10 | In progress | Design API",
    "Load testing | Ravi Patel | 2026-05-15 | Not started | Beta release",
    "",
    "DEPENDENCIES",
    "Dependency | Related Task | Status",
    "University SSO credentials | Design API | Missing",
    "Device farm for Android tests | Load testing | Requested",
    "",
    "RISKS",
    "Schedule risk: Design API is delayed, which threatens the 2026-03-15 mobile app deadline.",
    "Resource gap: No dedicated QA engineer is assigned.",
    "Unrealistic deadline: Beta release on 2026-05-01 depends on a delayed API design.",
    "Conflicting timeline: Load testing is scheduled after the beta release date but is listed as a quality gate.",
]


def write_college_docx(path: Path) -> None:
    document = Document()
    title = document.add_heading("College Project Plan", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    document.add_paragraph("PROJECT NAME: Smart Library Kiosk")
    document.add_paragraph(
        "OBJECTIVE: Create a self-service kiosk that lets students reserve books and print due-date slips."
    )
    document.add_paragraph(
        "SCOPE: Hardware kiosk UI, catalog search, and reservation workflow. Inventory purchasing is out of scope."
    )
    document.add_paragraph("START DATE: 2026-08-01")
    document.add_paragraph("DEADLINE: 2026-11-15")
    document.add_heading("Team members", level=1)
    document.add_paragraph("Maya Chen: Team lead")
    document.add_paragraph("Luis Romero: Frontend")
    document.add_paragraph("Ananya Iyer: Hardware integration")
    document.add_heading("Tasks", level=1)
    table = document.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "Task"
    hdr[1].text = "Owner"
    hdr[2].text = "Deadline"
    hdr[3].text = "Status"
    for row in [
        ("Wireframe kiosk screens", "Luis Romero", "2026-08-20", "Completed"),
        ("Catalog API prototype", "Maya Chen", "2026-09-10", "In progress"),
        ("Card reader integration", "Ananya Iyer", "2026-10-05", "Not started"),
    ]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    document.add_heading("Milestones", level=1)
    table = document.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "Milestone"
    hdr[1].text = "Date"
    hdr[2].text = "Owner"
    hdr[3].text = "Status"
    for row in [
        ("Proposal approved", "2026-08-10", "Maya Chen", "Completed"),
        ("Midterm demo", "2026-10-01", "Luis Romero", "Not started"),
        ("Final presentation", "2026-11-15", "Maya Chen", "Not started"),
    ]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    document.add_paragraph(
        "Note: Dependencies between hardware vendors are not listed in this plan."
    )
    document.save(path)


MEETING_NOTES = """\
MEETING NOTES - Smart Library Kiosk weekly standup
Date: 2026-09-12

BLOCKERS
Blocker: Card reader SDK license not approved | Owner: Ananya Iyer | Priority: High | Status: Open
Evidence: Procurement emailed that the SDK license is waiting on finance.

PENDING DECISIONS
Decision Required: Whether to use campus SSO or a local PIN | Responsible Person: Maya Chen | Due Date: 2026-09-18 | Status: Open
Evidence: Advisors asked the team to choose an authentication method before coding login.

OPEN ISSUES
Issue: Kiosk printer jams during due-date slip tests | Owner: Luis Romero | Priority: Medium | Status: Open
Evidence: Three of five print trials jammed in the lab on 2026-09-11.

ACTION ITEMS
Action Item: Send license reminder to finance | Owner: Ananya Iyer | Due Date: 2026-09-14 | Priority: High | Status: Open
Action Item: Draft SSO vs PIN comparison | Owner: Maya Chen | Due Date: 2026-09-17 | Priority: High | Status: Open
Action Item: Order spare printer rollers | Owner: Luis Romero | Due Date: 2026-09-20 | Priority: Medium | Status: Open
"""


def write_task_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tasks"
    sheet.append(["Task", "Owner", "Status", "Deadline", "Depends On"])
    for row in [
        ("Wireframe kiosk screens", "Luis Romero", "Completed", "2026-08-20", "Proposal approved"),
        ("Catalog API prototype", "Maya Chen", "In progress", "2026-09-10", "Wireframe kiosk screens"),
        ("Card reader integration", "Ananya Iyer", "Delayed", "2026-09-05", "Catalog API prototype"),
        ("Reservation flow tests", "Luis Romero", "Blocked", "2026-09-25", "Card reader integration"),
        ("User guide", "Maya Chen", "Not started", "2026-11-01", "Reservation flow tests"),
    ]:
        sheet.append(row)
    deps = workbook.create_sheet("Dependencies")
    deps.append(["Dependency", "Related Task", "Status"])
    deps.append(["Vendor SDK license", "Card reader integration", "Missing"])
    deps.append(["Lab kiosk hardware", "Reservation flow tests", "Requested"])
    workbook.save(path)


def write_status_docx(path: Path) -> None:
    document = Document()
    document.add_heading("Project Status Report", 0)
    document.add_paragraph("PROJECT NAME: Smart Library Kiosk")
    document.add_paragraph("Report date: 2026-09-12")
    document.add_paragraph("Current progress: Wireframes are complete. Catalog API prototype is in progress.")
    document.add_paragraph("Completed tasks: Wireframe kiosk screens, proposal approved.")
    document.add_paragraph("Delayed tasks: Card reader integration is delayed because the SDK license is missing.")
    document.add_heading("Risks", level=1)
    document.add_paragraph(
        "Schedule risk: Card reader integration slipped past 2026-09-05 and threatens the 2026-10-01 midterm demo."
    )
    document.add_paragraph("Resource gap: No dedicated tester is assigned.")
    document.add_heading("Blockers", level=1)
    document.add_paragraph(
        "Blocker: Card reader SDK license not approved | Owner: Ananya Iyer | Priority: High | Status: Open"
    )
    document.add_heading("Upcoming milestones", level=1)
    table = document.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "Milestone"
    hdr[1].text = "Date"
    hdr[2].text = "Owner"
    hdr[3].text = "Status"
    for row in [
        ("Midterm demo", "2026-10-01", "Luis Romero", "At risk"),
        ("Final presentation", "2026-11-15", "Maya Chen", "Not started"),
    ]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    document.save(path)


def main() -> None:
    write_text_pdf(ROOT / "software_development_project_plan.pdf", SOFTWARE_PLAN)
    write_college_docx(ROOT / "college_project_plan.docx")
    (ROOT / "meeting_notes.txt").write_text(MEETING_NOTES, encoding="utf-8")
    write_task_xlsx(ROOT / "project_task_tracker.xlsx")
    write_status_docx(ROOT / "project_status_report.docx")
    print(f"Wrote sample documents to {ROOT}")


if __name__ == "__main__":
    main()
