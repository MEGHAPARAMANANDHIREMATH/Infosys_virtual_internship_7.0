from pathlib import Path

from docx import Document as DocxDocument

from ingestion.exceptions import ParsingError
from ingestion.parsers.base import ParsedSegment


class DocxParser:
    def parse(self, file_path: Path) -> list[ParsedSegment]:
        try:
            document = DocxDocument(str(file_path))
            segments: list[ParsedSegment] = []
            current_section = ""

            for paragraph in document.paragraphs:
                style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
                text = paragraph.text or ""
                if "heading" in style_name and text.strip():
                    current_section = text.strip()
                segments.append(
                    ParsedSegment(
                        text=text,
                        section=current_section or None,
                        extra={"source_type": "docx", "style": style_name},
                    )
                )

            for table_index, table in enumerate(document.tables, start=1):
                rows: list[str] = []
                for row in table.rows:
                    cells = [(cell.text or "").strip() for cell in row.cells]
                    if any(cells):
                        rows.append(" | ".join(cells))
                if rows:
                    segments.append(
                        ParsedSegment(
                            text="Table:\n" + "\n".join(rows),
                            section=current_section or f"table_{table_index}",
                            extra={"source_type": "docx_table", "table_index": table_index},
                        )
                    )
            return segments
        except Exception as exc:
            raise ParsingError(f"DOCX parsing failed: {exc}") from exc
